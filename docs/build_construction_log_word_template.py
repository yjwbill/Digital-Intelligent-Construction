from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\loulidong\Desktop\新建文件夹\Digital-Intelligent-Construction\docs\施工日志在线填报导出模板.docx"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LABEL_FILL = "F2F4F7"
BORDER = "C9D1DB"
MUTED = RGBColor(93, 105, 120)
TOTAL_DXA = 14400

def set_font(run, size=9, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(str(text)), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_table_geometry(table, widths):
    assert sum(widths) == TOTAL_DXA
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TOTAL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")

def style_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), BORDER)
        borders.append(el)

def add_section_heading(doc, title, suffix=""):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(title), size=12, bold=True, color=RGBColor(46, 116, 181))
    if suffix:
        set_font(p.add_run(f"  {suffix}"), size=9, color=MUTED)

def add_key_value_grid(doc, items):
    table = doc.add_table(rows=0, cols=8)
    table.style = "Table Grid"
    widths = [1150, 2450] * 4
    for start in range(0, len(items), 4):
        row = table.add_row()
        chunk = items[start:start+4]
        for idx in range(4):
            label, value = chunk[idx] if idx < len(chunk) else ("", "")
            set_cell_text(row.cells[idx*2], label, bold=True, size=8, color=RGBColor(70, 82, 98))
            shade(row.cells[idx*2], LABEL_FILL)
            set_cell_text(row.cells[idx*2+1], value, size=8.5)
    set_table_geometry(table, widths)
    style_table_borders(table)
    return table

def add_data_table(doc, headers, rows, widths, sizes=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=7.8, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(45, 61, 80))
        shade(table.rows[0].cells[i], LIGHT_BLUE)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) or len(str(value)) < 10 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], value, size=(sizes[i] if sizes else 8), align=align)
    set_table_geometry(table, widths)
    style_table_borders(table)
    return table

def add_detail_block(doc, title, items):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(title), size=9, bold=True, color=RGBColor(31, 77, 120))
    add_key_value_grid(doc, items)

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.25
for name, size, before, after, color in (("Heading 1", 16, 18, 10, BLUE), ("Heading 2", 13, 14, 7, BLUE), ("Heading 3", 12, 10, 5, "1F4D78")):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(hp.add_run("数智施工 · 施工日志导出"), size=8, color=MUTED)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("施工日志归档文件  |  系统导出"), size=8, color=MUTED)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(2)
set_font(title.add_run("施 工 日 志"), size=18, bold=True, color=RGBColor(31, 77, 120))
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(8)
set_font(subtitle.add_run("在线填报导出版"), size=9, color=MUTED)

add_section_heading(doc, "基础信息")
add_key_value_grid(doc, [
    ("项目名称", "机场联络线工程"), ("施工工区", "主体结构区、附属结构 A 区"),
    ("日期", "2026-08-04"), ("星期", "星期二"),
    ("温度", "25℃"), ("天气是否影响工作", "否"),
    ("记录人", "楼力栋、张三"), ("上报时间", "2026-08-04 18:20")
])

add_section_heading(doc, "人员信息")
add_key_value_grid(doc, [
    ("总包管理人员", "21 人"), ("分包管理人员", "12 人"), ("劳务人员", "39 人")
])

add_section_heading(doc, "今日主要工作", "2026-08-04")
today_headers = ["序号", "施工工区", "施工分项", "施工部位", "工作内容", "工作进度", "施工图片", "记录人", "备注"]
today_widths = [500, 1350, 1550, 1450, 3200, 1750, 1500, 950, 2150]
add_data_table(doc, today_headers, [
    ["1", "主体结构区", "主体结构施工", "地下二层中板", "完成钢筋绑扎、模板复核及隐蔽工程验收。", "按计划完成", "施工图片 1-2", "楼力栋", "材料、机具及安全防护检查正常。"],
    ["2", "附属结构 A 区", "土方开挖", "1 号风亭", "完成土方开挖 320m³，同步实施临边防护。", "完成 85%", "施工图片 3-4", "张三", "监测数据正常。"]
], today_widths)

add_section_heading(doc, "明日主要工作", "2026-08-05")
tomorrow_headers = ["序号", "施工工区", "施工分项", "施工部位", "工作内容", "施工图片", "记录人", "备注"]
tomorrow_widths = [550, 1500, 1750, 1600, 3500, 1700, 1000, 2800]
add_data_table(doc, tomorrow_headers, [
    ["1", "主体结构区", "主体结构施工", "地下二层中板", "继续开展混凝土浇筑准备及现场安全巡查。", "计划照片", "楼力栋", "提前落实材料进场计划。"],
    ["2", "附属结构 A 区", "土方开挖", "1 号风亭", "完成剩余土方开挖并组织基底验收。", "计划照片", "张三", "关注降雨和基坑排水。"]
], tomorrow_widths)

doc.add_page_break()
add_section_heading(doc, "里程碑节点情况", "进行中")
add_detail_block(doc, "节点 1", [
    ("里程碑节点名称", "主体结构封顶"), ("计划完成日期（最新）", "2026-08-20"),
    ("节点状态", "延期"), ("管控等级", "子公司管控"),
    ("是否重点进度节点", "是"), ("里程碑情况", "进度可控"),
    ("里程碑进展情况", "现场资源配置已完成，按调整计划持续推进。")
])

add_section_heading(doc, "里程碑节点情况", "已完成")
add_detail_block(doc, "节点 1", [
    ("里程碑节点名称", "围护结构完成"), ("计划完成日期（最新）", "2026-07-10"),
    ("节点状态", "已完成"), ("管控等级", "项目管控"),
    ("是否重点进度节点", "否"), ("实际完成日期", "2026-07-08")
])

add_section_heading(doc, "风险情况")
add_detail_block(doc, "风险 1", [
    ("风险类型", "深基坑开挖"), ("风险名称", "附属结构土方开挖"),
    ("风险等级", "II 级"), ("是否完成", "否"),
    ("是否受控", "是"), ("风险情况", "风险可控"),
    ("风险进展情况", "现场监测数据正常，按专项方案组织施工。")
])
add_detail_block(doc, "风险 2", [
    ("风险类型", "承重支模架"), ("风险名称", "主体结构模板支撑"),
    ("风险等级", "II 级"), ("是否完成", "否"),
    ("是否受控", "是"), ("风险情况", "风险可控"),
    ("风险进展情况", "验收记录齐全，持续开展班前检查。")
])

add_section_heading(doc, "发生停工情况")
stop_table = doc.add_table(rows=1, cols=1)
set_cell_text(stop_table.cell(0, 0), "无停工情况。", size=9)
set_table_geometry(stop_table, [TOTAL_DXA])
style_table_borders(stop_table)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(8)
note.paragraph_format.space_after = Pt(0)
set_font(note.add_run("说明：本文件结构与在线填报页面一致；实际导出时由系统写入真实数据，并将施工图片嵌入对应单元格。"), size=8, color=MUTED)

doc.core_properties.title = "施工日志在线填报导出模板"
doc.core_properties.subject = "数智施工施工日志 Word 导出样稿"
doc.core_properties.author = "数智施工"
doc.save(OUT)
print(OUT)
